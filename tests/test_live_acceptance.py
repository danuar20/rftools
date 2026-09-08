"""
Automated Acceptance Tests for Acceptance Criteria AC 1 to 7
Tests end-to-end processing across all 6 RF & GIS tools.
"""

import io
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
import pytest
from starlette.testclient import TestClient

from backend.main import app
from backend.tools.template_generator import generate_template

client = TestClient(app)

def test_ac1_startup_and_health():
    """AC 1: App startup, health endpoint, and all 6 tools active."""
    res = client.get("/api/v1/health")
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "healthy"
    assert data["port"] == 5005
    engines = data["engines"]
    for eng in ["excel_to_kml", "prb_kml", "isd_calculator", "geohash_to_shp", "geohash_to_latlon", "latlon_to_geohash"]:
        assert engines[eng] == "active"

    res_tools = client.get("/api/v1/tools")
    assert res_tools.status_code == 200
    tools = res_tools.json()
    assert len(tools) == 6

def test_ac2_point_kml():
    """AC 2: Point KML Placemark generation."""
    xlsx_bytes, filename = generate_template("point_kml")
    files = {"file": (filename, xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    data = {"scale": "1.25", "color_rgb": "#0EA5E9", "label_color": "#FFFF00"}

    res = client.post("/api/v1/tools/excel-to-kml/process", files=files, data=data)
    assert res.status_code == 200
    assert "kml" in res.headers["content-type"].lower()

    root = ET.fromstring(res.content.decode("utf-8"))
    for elem in root.iter():
        if "}" in elem.tag:
            elem.tag = elem.tag.split("}", 1)[1]

    placemarks = root.findall(".//Placemark")
    assert len(placemarks) == 4
    names = [p.find("name").text for p in placemarks if p.find("name") is not None]
    assert any("JKT001" in n for n in names)

def test_ac3_prb_kml():
    """AC 3: PRB 3D extruded sectors at carrier altitude with color thresholding."""
    xlsx_bytes, filename = generate_template("prb_kml")
    files = {"file": (filename, xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    data = {"color_by_metric": "DL_PRB", "opacity_percent": "40"}

    res = client.post("/api/v1/tools/prb-kml/process", files=files, data=data)
    assert res.status_code == 200
    assert "kml" in res.headers["content-type"].lower()

    root = ET.fromstring(res.content.decode("utf-8"))
    for elem in root.iter():
        if "}" in elem.tag:
            elem.tag = elem.tag.split("}", 1)[1]

    polygons = root.findall(".//Polygon")
    assert len(polygons) == 3
    for poly in polygons:
        assert poly.find("altitudeMode").text == "relativeToGround"
        assert poly.find("extrude").text == "1"

    coords = polygons[0].find(".//coordinates").text.strip().split()
    assert float(coords[0].split(",")[2]) == 38.0

def test_ac4_isd_calculator():
    """AC 4: ISD calculation with n_nearest=2 and Summary sheet."""
    bytes_a, fn_a = generate_template("isd_a")
    bytes_b, fn_b = generate_template("isd_b")
    files = {
        "file_a": (fn_a, bytes_a, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "file_b": (fn_b, bytes_b, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    data = {"n_nearest": "2"}

    res = client.post("/api/v1/tools/isd-calculator/process", files=files, data=data)
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "ISD_Result" in wb.sheetnames
    assert "Summary" in wb.sheetnames

    ws_res = wb["ISD_Result"]
    headers = [cell.value for cell in ws_res[1]]
    assert "distance_km" in headers
    rows = list(ws_res.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 4

    ws_sum = wb["Summary"]
    sum_dict = dict(list(ws_sum.iter_rows(values_only=True)))
    assert sum_dict["count"] == 4
    assert sum_dict["min_km"] <= sum_dict["max_km"]

def test_ac5_shapefile_generator():
    """AC 5: Shapefile zip package with DBF headers <= 10 chars."""
    xlsx_bytes, filename = generate_template("geohash")
    files = {"file": (filename, xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    data = {"mode": "custom", "size_m": "500.0", "geohash_col": "GEOHASH"}

    res = client.post("/api/v1/tools/geohash-to-shp/process", files=files, data=data)
    assert res.status_code == 200
    assert "zip" in res.headers["content-type"]

    z = zipfile.ZipFile(io.BytesIO(res.content))
    names = z.namelist()
    for ext in ["shp", "shx", "dbf", "prj"]:
        assert any(n.endswith(f".{ext}") for n in names)

    dbf_file = [n for n in names if n.endswith(".dbf")][0]
    dbf_bytes = z.read(dbf_file)
    offset = 32
    while offset < len(dbf_bytes) and dbf_bytes[offset] != 0x0D:
        fname = dbf_bytes[offset:offset+11].split(b"\x00")[0].decode("ascii", errors="ignore")
        if fname:
            assert len(fname) <= 10
        offset += 32

def test_ac6_geohash_roundtrip():
    """AC 6: Geohash roundtrip fidelity within +-0.00078 degrees."""
    xlsx_bytes, filename = generate_template("latlon")
    files_enc = {"file": (filename, xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    data_enc = {"lat_col": "LAT", "lon_col": "LONG", "precision": "7", "output_format": "xlsx"}

    res_enc = client.post("/api/v1/tools/latlon-to-geohash/process", files=files_enc, data=data_enc)
    assert res_enc.status_code == 200

    wb_enc = openpyxl.load_workbook(io.BytesIO(res_enc.content))
    ws_enc = wb_enc.active
    headers = [c.value for c in ws_enc[1]]
    gh_idx = headers.index("geohash")
    gh_val = ws_enc.cell(row=2, column=gh_idx+1).value
    assert gh_val.startswith("qqgu")

    files_dec = {"file": ("enc.xlsx", res_enc.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    data_dec = {"geohash_col": "geohash", "output_format": "xlsx"}
    res_dec = client.post("/api/v1/tools/geohash-to-latlon/process", files=files_dec, data=data_dec)
    assert res_dec.status_code == 200

    wb_dec = openpyxl.load_workbook(io.BytesIO(res_dec.content))
    ws_dec = wb_dec.active
    dec_headers = [c.value for c in ws_dec[1]]
    lat_idx = dec_headers.index("latitude")
    lon_idx = dec_headers.index("longitude")
    dec_lat = float(ws_dec.cell(row=2, column=lat_idx+1).value)
    dec_lon = float(ws_dec.cell(row=2, column=lon_idx+1).value)

    assert abs(dec_lat - (-6.175392)) <= 0.00078
    assert abs(dec_lon - 106.827153) <= 0.00078

def test_ac7_templates_download():
    """AC 7: Sample template generation and download."""
    for tid in ["point_kml", "prb_kml", "isd_a", "isd_b", "geohash", "latlon"]:
        res = client.get(f"/api/v1/templates/{tid}")
        assert res.status_code == 200
        assert "spreadsheetml" in res.headers["content-type"]
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        assert "Sheet1" in wb.sheetnames
        ws = wb["Sheet1"]
        assert len(list(ws.iter_rows(values_only=True))) >= 2
