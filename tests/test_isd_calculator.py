"""
Tests for Tool 3: ISD (Inter-Site Distance) Calculator
Verifies Haversine formula calculation, N-nearest neighbor ranking, and summary stats.
"""

import io
import openpyxl
import pytest
from backend.tools.isd_calculator import haversine_km, calculate_isd

def test_haversine_formula_accuracy():
    # Monas to GBK (~8.05 km)
    monas_lat, monas_lon = -6.175392, 106.827153
    gbk_lat, gbk_lon = -6.243589, 106.801444
    d = haversine_km(monas_lat, monas_lon, gbk_lat, gbk_lon)
    assert 7.9 < d < 8.2
    # Zero distance for identical points
    assert haversine_km(monas_lat, monas_lon, monas_lat, monas_lon) == 0.0

def create_isd_file_a():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["SITE_ID", "LAT", "LONG"])
    ws.append(["SITE_A01", -6.175392, 106.827153]) # Monas
    ws.append(["SITE_A02", -6.917464, 107.619123]) # Bandung Gedung Sate
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()

def create_isd_file_b():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["SITE_ID", "LAT", "LONG"])
    ws.append(["SITE_B01", -6.243589, 106.801444]) # GBK (~8km from Monas)
    ws.append(["SITE_B02", -6.183333, 106.841667]) # Gambir area (~1.8km from Monas)
    ws.append(["SITE_B03", -6.921852, 107.607140]) # Bandung Alun-Alun (~1.4km from Gedung Sate)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()

def test_calculate_isd_multi_neighbor():
    file_a = create_isd_file_a()
    file_b = create_isd_file_b()

    xlsx_bytes, summary = calculate_isd(file_a, file_b, n_nearest=2)
    assert isinstance(xlsx_bytes, bytes)
    assert len(xlsx_bytes) > 0

    assert summary["sites_a_count"] == 2
    assert summary["sites_b_count"] == 3
    assert summary["n_nearest"] == 2
    assert summary["total_pairs"] == 4 # 2 sites * 2 nearest = 4 pairs
    assert summary["min_km"] > 0
    assert summary["max_km"] > summary["min_km"]

    # Verify produced workbook structure
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "ISD_Result" in wb.sheetnames
    assert "Summary" in wb.sheetnames

    res_sheet = wb["ISD_Result"]
    rows = list(res_sheet.iter_rows(values_only=True))
    assert rows[0] == ("siteA", "latA", "lonA", "nearestSiteB", "latB", "lonB", "distance_km")
    # 1 header + 4 data rows
    assert len(rows) == 5

    # Check nearest neighbor for SITE_A01: Gambir (B02) should be closer than GBK (B01)
    assert rows[1][0] == "SITE_A01"
    assert rows[1][3] == "SITE_B02"
    assert rows[2][0] == "SITE_A01"
    assert rows[2][3] == "SITE_B01"
    assert rows[1][6] < rows[2][6]
