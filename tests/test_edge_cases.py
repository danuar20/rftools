"""
Edge Case & Error Handling Tests for RF TOOLS Calculation Engines & API
"""

import io
import openpyxl
import pytest
from fastapi.testclient import TestClient
from backend.main import app
from backend.tools.constants import get_band_params, get_dl_prb_color, get_ul_prb_color, get_rrc_color
from backend.tools.excel_to_kml import convert_excel_to_kml
from backend.tools.prb_kml import convert_excel_to_prb_kml
from backend.tools.isd_calculator import calculate_isd
from backend.tools.geohash_to_shp import convert_geohash_to_shp
from backend.tools.column_inspector import classify_column, inspect_file_columns

client = TestClient(app)

def test_classify_column_edge_cases():
    assert classify_column("Latitude", [])["role"] == "latitude"
    assert classify_column("LAT", [])["role"] == "latitude"
    assert classify_column("longitude_deg", [])["role"] == "longitude"
    assert classify_column("GEOHASH", [])["role"] == "geohash"
    assert classify_column("bearing", [])["role"] == "azimuth"
    assert classify_column("carrier", [])["role"] == "beam"
    assert classify_column("DL_PRB_AVG", [])["role"] == "dl_prb"
    assert classify_column("RRC_USERS", [])["role"] == "rrc_user"
    assert classify_column("unknown_col_xyz", [])["role"] is None

def test_extreme_geohash_precisions():
    from backend.tools.latlon_to_geohash import convert_latlon_to_geohash
    csv_text = "lat,lon\n-6.175392,106.827153\n"
    # Precision clamped to 1
    _, s1 = convert_latlon_to_geohash(csv_text.encode(), filename="t.csv", precision=0, output_format="csv")
    assert s1["precision"] == 1
    # Precision clamped to 12
    _, s12 = convert_latlon_to_geohash(csv_text.encode(), filename="t.csv", precision=15, output_format="csv")
    assert s12["precision"] == 12

def test_prb_kml_missing_required_columns():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["LAT", "LONG"]) # missing BEAM and DIRECTION
    ws.append([-6.175, 106.827])
    bio = io.BytesIO()
    wb.save(bio)
    with pytest.raises(ValueError, match="Missing required columns"):
        convert_excel_to_prb_kml(bio.getvalue())

def test_isd_single_site_vs_empty():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["SITE_ID", "LAT", "LONG"])
    bio_empty = io.BytesIO()
    wb.save(bio_empty)

    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.append(["SITE_ID", "LAT", "LONG"])
    ws2.append(["A1", -6.1, 106.8])
    bio_a = io.BytesIO()
    wb2.save(bio_a)

    with pytest.raises(ValueError, match="empty or missing"):
        calculate_isd(bio_empty.getvalue(), bio_a.getvalue())

def test_api_404_tool():
    res = client.get("/api/v1/tools/non-existent-tool")
    assert res.status_code == 404

def test_api_404_template():
    res = client.get("/api/v1/templates/invalid_template_id")
    assert res.status_code == 404

def test_api_preview_mode_excel_to_kml():
    from backend.tools.template_generator import generate_template
    wb_bytes, _ = generate_template("point_kml")
    files = {"file": ("points.xlsx", wb_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res = client.post("/api/v1/tools/excel-to-kml/process?preview=true", files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert "summary" in data
    assert data["summary"]["valid_points"] > 0
