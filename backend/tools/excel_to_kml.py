"""
Tool 1: Excel to Point KML Converter
Extracts coordinates from Excel sheets and generates Google Earth placemarks.
"""

import io
from typing import Optional, Tuple, Dict, Any, List
import openpyxl
import simplekml
from backend.tools.constants import (
    DEFAULT_POINT_ICON,
    DEFAULT_POINT_COLOR,
    DEFAULT_LABEL_COLOR,
    DEFAULT_POINT_SCALE,
)

def parse_color(c) -> Tuple[int, int, int]:
    """Parse RGB tuple, list, or hex string into (r, g, b)."""
    if isinstance(c, (list, tuple)) and len(c) == 3:
        return (int(c[0]), int(c[1]), int(c[2]))
    if isinstance(c, str):
        c = c.lstrip('#')
        if len(c) == 6:
            return (int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16))
    return DEFAULT_POINT_COLOR

def find_latlon_columns(headers: List[str]) -> Tuple[Optional[int], Optional[int]]:
    """Find column indices for latitude and longitude from headers list."""
    lat_idx, lon_idx = None, None
    low = [str(h).strip().lower() if h is not None else "" for h in headers]
    for i, h in enumerate(low):
        if h in ('lat', 'latitude', 'y'):
            lat_idx = i
        elif h in ('lon', 'long', 'longitude', 'x'):
            lon_idx = i
    return lat_idx, lon_idx

def convert_excel_to_kml(
    file_bytes_or_path,
    folder_name: str = "SITENAME",
    icon_url: Optional[str] = None,
    color_rgb: Any = DEFAULT_POINT_COLOR,
    label_color: Any = DEFAULT_LABEL_COLOR,
    scale: float = DEFAULT_POINT_SCALE,
    lat_col: Optional[str] = None,
    lon_col: Optional[str] = None,
    name_col: Optional[str] = None,
    sheet_name: Optional[str] = None
) -> Tuple[bytes, Dict[str, Any]]:
    """
    Converts Excel file to KML placemark points.
    Returns (kml_bytes, result_summary).
    """
    if isinstance(file_bytes_or_path, (bytes, bytearray)):
        wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes_or_path), data_only=True)
    elif hasattr(file_bytes_or_path, 'read'):
        wb = openpyxl.load_workbook(filename=file_bytes_or_path, data_only=True)
    else:
        wb = openpyxl.load_workbook(filename=file_bytes_or_path, data_only=True)

    if sheet_name and sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
    else:
        sheet = wb.active

    rows = list(sheet.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        raise ValueError("Excel file is empty or contains no data rows.")

    headers = [str(h) if h is not None else f"Col_{i}" for i, h in enumerate(rows[0])]
    header_map = {h.strip().lower(): i for i, h in enumerate(headers)}

    lat_idx = header_map.get(lat_col.strip().lower()) if lat_col else None
    lon_idx = header_map.get(lon_col.strip().lower()) if lon_col else None
    name_idx = header_map.get(name_col.strip().lower()) if name_col else None

    # Auto-detect if not specified
    if lat_idx is None or lon_idx is None:
        det_lat, det_lon = find_latlon_columns(headers)
        lat_idx = det_lat if lat_idx is None else lat_idx
        lon_idx = det_lon if lon_idx is None else lon_idx

    # Fallback to column index heuristics
    if lat_idx is None or lon_idx is None:
        if len(headers) >= 4:
            lon_idx = 2 if lon_idx is None else lon_idx
            lat_idx = 3 if lat_idx is None else lat_idx

    if lat_idx is None or lon_idx is None:
        raise ValueError("Could not identify Latitude and Longitude columns in the file.")

    # Name column heuristics
    if name_idx is None:
        for candidate in ['site', 'sitename', 'site_name', 'site_id', 'siteid', 'name']:
            if candidate in header_map:
                name_idx = header_map[candidate]
                break
        if name_idx is None and len(headers) >= 4:
            name_idx = 4 if len(headers) > 4 else 1

    rgb = parse_color(color_rgb)
    lbl_rgb = parse_color(label_color)

    k = simplekml.Kml()
    fol = k.newfolder(name=folder_name or "SITENAME")

    style = simplekml.Style()
    style.iconstyle.icon = simplekml.Icon(href=icon_url or DEFAULT_POINT_ICON)
    style.iconstyle.color = simplekml.Color.rgb(*rgb)
    style.iconstyle.scale = scale
    style.labelstyle.color = simplekml.Color.rgb(*lbl_rgb)
    style.labelstyle.scale = scale

    data_rows = rows[1:]
    total_rows = len(data_rows)
    valid_points = 0
    skipped_points = 0
    preview_rows = []

    for r in data_rows:
        try:
            if not r or all(c is None for c in r):
                continue
            lat_val = r[lat_idx] if lat_idx < len(r) else None
            lon_val = r[lon_idx] if lon_idx < len(r) else None

            if lat_val is None or lon_val is None or str(lat_val).strip() == '' or str(lon_val).strip() == '':
                skipped_points += 1
                continue

            lat_f = float(lat_val)
            lon_f = float(lon_val)

            if not (-90.0 <= lat_f <= 90.0 and -180.0 <= lon_f <= 180.0):
                skipped_points += 1
                continue

            name = None
            if name_idx is not None and name_idx < len(r) and r[name_idx] not in (None, ""):
                name = str(r[name_idx]).strip()
            else:
                for cand in ['sitename', 'site_name', 'site_id', 'siteid', 'name', 'site']:
                    if cand in header_map:
                        idx_cand = header_map[cand]
                        if idx_cand not in (lat_idx, lon_idx) and idx_cand < len(r) and r[idx_cand] not in (None, ""):
                            name = str(r[idx_cand]).strip()
                            break
                if not name:
                    for try_idx in range(len(r)):
                        if try_idx not in (lat_idx, lon_idx) and r[try_idx] not in (None, ""):
                            name = str(r[try_idx]).strip()
                            break
            final_name = name or f"Site_{valid_points + 1}"

            p = fol.newpoint(name=final_name, coords=[(lon_f, lat_f)])
            p.style = style
            valid_points += 1

            if len(preview_rows) < 10:
                preview_rows.append({
                    "name": final_name,
                    "latitude": lat_f,
                    "longitude": lon_f
                })
        except Exception:
            skipped_points += 1
            continue

    if valid_points == 0:
        raise ValueError("No valid coordinate rows could be converted.")

    kml_str = k.kml()
    summary = {
        "tool": "excel_to_kml",
        "folder_name": folder_name,
        "total_rows": total_rows,
        "valid_points": valid_points,
        "skipped_points": skipped_points,
        "preview_rows": preview_rows,
    }
    return kml_str.encode('utf-8'), summary
