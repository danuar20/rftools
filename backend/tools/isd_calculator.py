"""
Tool 3: ISD (Inter-Site Distance) Calculator
Computes Haversine great-circle distances between source sites (File A)
and candidate sites (File B) finding N-nearest neighbors.
Supports both kilometers (km) and meters (m) units.
"""

import io
import math
from typing import Optional, Dict, Any, Tuple, List
import openpyxl
from backend.tools.constants import EARTH_RADIUS_HAVERSINE_KM
from backend.tools.excel_to_kml import find_latlon_columns

def haversine_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    """Compute great-circle distance between two points (in degrees) in kilometers."""
    rlat1 = math.radians(lat1)
    rlon1 = math.radians(lon1)
    rlat2 = math.radians(lat2)
    rlon2 = math.radians(lon2)
    dlat = rlat2 - rlat1
    dlon = rlon2 - rlon1
    a = math.sin(dlat / 2.0)**2 + math.cos(rlat1) * math.cos(rlat2) * math.sin(dlon / 2.0)**2
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1.0 - a))
    return EARTH_RADIUS_HAVERSINE_KM * c

def _load_sites_from_excel(file_bytes_or_path, lat_col=None, lon_col=None, name_col=None):
    if isinstance(file_bytes_or_path, (bytes, bytearray)):
        wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes_or_path), data_only=True)
    elif hasattr(file_bytes_or_path, 'read'):
        wb = openpyxl.load_workbook(filename=file_bytes_or_path, data_only=True)
    else:
        wb = openpyxl.load_workbook(filename=file_bytes_or_path, data_only=True)

    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        raise ValueError("File is empty or missing data rows.")

    headers = [str(h).strip() if h is not None else f"Col_{i}" for i, h in enumerate(rows[0])]
    header_map = {h.lower(): i for i, h in enumerate(headers)}

    lat_idx = header_map.get(lat_col.strip().lower()) if lat_col else None
    lon_idx = header_map.get(lon_col.strip().lower()) if lon_col else None
    name_idx = header_map.get(name_col.strip().lower()) if name_col else None

    if lat_idx is None or lon_idx is None:
        det_lat, det_lon = find_latlon_columns(headers)
        lat_idx = det_lat if lat_idx is None else lat_idx
        lon_idx = det_lon if lon_idx is None else lon_idx

    if lat_idx is None or lon_idx is None:
        if len(headers) >= 4:
            lon_idx = 2 if lon_idx is None else lon_idx
            lat_idx = 3 if lat_idx is None else lat_idx

    if lat_idx is None or lon_idx is None:
        raise ValueError("Could not detect Latitude/Longitude columns.")

    if name_idx is None:
        for candidate in ['site_id', 'siteid', 'sitename', 'site', 'name', 'id']:
            if candidate in header_map:
                name_idx = header_map[candidate]
                break
        if name_idx is None:
            name_idx = 0

    sites = []
    for r in rows[1:]:
        try:
            if not r or all(c is None for c in r):
                continue
            lat_val = r[lat_idx] if lat_idx < len(r) else None
            lon_val = r[lon_idx] if lon_idx < len(r) else None
            if lat_val is None or lon_val is None:
                continue
            lat = float(lat_val)
            lon = float(lon_val)
            if not (-90 <= lat <= 90 and -180 <= lon <= 180):
                continue
            name = None
            if name_idx < len(r) and r[name_idx] not in (None, ""):
                name = str(r[name_idx]).strip()
            else:
                for try_idx in (0, 1, 4):
                    if try_idx < len(r) and r[try_idx] not in (None, ""):
                        name = str(r[try_idx]).strip()
                        break
            sites.append((name or f"Site_{len(sites)+1}", lat, lon))
        except Exception:
            continue

    return sites

def calculate_isd(
    file_a_bytes_or_path,
    file_b_bytes_or_path,
    n_nearest: int = 1,
    distance_unit: str = "km",
    lat_col_a: Optional[str] = None,
    lon_col_a: Optional[str] = None,
    name_col_a: Optional[str] = None,
    lat_col_b: Optional[str] = None,
    lon_col_b: Optional[str] = None,
    name_col_b: Optional[str] = None
) -> Tuple[bytes, Dict[str, Any]]:
    """
    Computes nearest neighbor distances between File A and File B.
    Supports distance_unit='km' (default) or 'm' (meters).
    Returns (excel_bytes, summary_dict).
    """
    n_nearest = max(1, min(5, int(n_nearest)))
    is_meters = str(distance_unit).strip().lower() in ("m", "meter", "meters")
    dist_col_name = "distance_m" if is_meters else "distance_km"
    multiplier = 1000.0 if is_meters else 1.0
    decimals = 2 if is_meters else 6

    sites_a = _load_sites_from_excel(file_a_bytes_or_path, lat_col_a, lon_col_a, name_col_a)
    if not sites_a:
        raise ValueError("No valid coordinate rows found in File A.")

    sites_b = _load_sites_from_excel(file_b_bytes_or_path, lat_col_b, lon_col_b, name_col_b)
    if not sites_b:
        raise ValueError("No valid coordinate rows found in File B.")

    out_wb = openpyxl.Workbook()
    out_sh = out_wb.active
    out_sh.title = "ISD_Result"
    out_sh.append(["siteA", "latA", "lonA", "nearestSiteB", "latB", "lonB", dist_col_name])

    distances_km = []
    distances_output = []
    preview_rows = []

    for name_a, lat_a, lon_a in sites_a:
        dist_list = []
        for name_b, lat_b, lon_b in sites_b:
            d_km = haversine_km(lat_a, lon_a, lat_b, lon_b)
            dist_list.append(((name_b, lat_b, lon_b), d_km))

        dist_list.sort(key=lambda x: x[1])

        k = min(n_nearest, len(dist_list))
        for i in range(k):
            (name_b, lat_b, lon_b), dist_km = dist_list[i]
            d_out = round(dist_km * multiplier, decimals)
            out_sh.append([
                name_a, lat_a, lon_a,
                name_b, lat_b, lon_b,
                d_out
            ])
            distances_km.append(dist_km)
            distances_output.append(d_out)

            if len(preview_rows) < 10:
                preview_row = {
                    "siteA": name_a,
                    "latA": lat_a,
                    "lonA": lon_a,
                    "nearestSiteB": name_b,
                    "latB": lat_b,
                    "lonB": lon_b,
                    "distance": d_out,
                    "unit": "m" if is_meters else "km",
                }
                preview_row[dist_col_name] = d_out
                preview_rows.append(preview_row)

    summary_sh = out_wb.create_sheet(title="Summary")
    if distances_output:
        min_val = round(min(distances_output), decimals)
        max_val = round(max(distances_output), decimals)
        mean_val = round(sum(distances_output) / len(distances_output), decimals)

        min_km = round(min(distances_km), 6)
        max_km = round(max(distances_km), 6)
        mean_km = round(sum(distances_km) / len(distances_km), 6)

        summary_sh.append(["count", len(distances_output)])
        summary_sh.append([f"min_{'m' if is_meters else 'km'}", min_val])
        summary_sh.append([f"mean_{'m' if is_meters else 'km'}", mean_val])
        summary_sh.append([f"max_{'m' if is_meters else 'km'}", max_val])
    else:
        min_val = max_val = mean_val = 0.0
        min_km = max_km = mean_km = 0.0
        summary_sh.append(["No valid distances computed"])

    output_buffer = io.BytesIO()
    out_wb.save(output_buffer)
    xlsx_bytes = output_buffer.getvalue()

    summary = {
        "tool": "isd_calculator",
        "sites_a_count": len(sites_a),
        "sites_b_count": len(sites_b),
        "n_nearest": n_nearest,
        "distance_unit": "m" if is_meters else "km",
        "total_pairs": len(distances_output),
        "min_distance": min_val,
        "mean_distance": mean_val,
        "max_distance": max_val,
        "min_km": min_km,
        "mean_km": mean_km,
        "max_km": max_km,
        "preview_rows": preview_rows,
    }
    return xlsx_bytes, summary
